"""Put super_img2ppt slides into the competition deck as native, editable slides.

An img2ppt build maps its canvas onto a 960 pt-wide slide; a 1536×864 scene is therefore exactly one 16:9 slide,
so its shapes are copied 1:1 (no scaling) onto a new slide that uses the deck's own title bar (title placeholder +
the two blue rectangles every content slide carries). Pictures (raster art kept as independent assets) are carried
over with their image parts.

usage: deck_merge.py --deck deck.pptx --out out.pptx [--font 微软雅黑] [--after 9] \
           <build/editable.pptx>::<slide title> [<build2/editable.pptx>::<title 2> ...]
"""
import argparse
import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
SLIDE_W_EMU = 12192000


def title_like(model_slide):
    """The deck's content slides share one title style; take it from a model slide (slide 9)."""
    return model_slide.shapes.title


def set_title(el, text):
    """Rewrite the title runs of a <p:sp> element: one run, the model's run properties."""
    txBody = el.find('p:txBody', NS)
    paras = txBody.findall('a:p', NS)
    p0 = paras[0]
    runs = p0.findall('a:r', NS)
    rpr = copy.deepcopy(runs[0].find('a:rPr', NS))
    for r in runs:
        p0.remove(r)
    for extra in paras[1:]:
        txBody.remove(extra)
    r = p0.makeelement('{%s}r' % NS['a'], {})
    r.append(rpr)
    t = r.makeelement('{%s}t' % NS['a'], {})
    t.text = text
    r.append(t)
    end = p0.find('a:endParaRPr', NS)
    if end is not None:
        end.addprevious(r)
    else:
        p0.append(r)


def copy_shape(el, dst_spTree, next_id):
    """Deep-copy one spTree child onto the destination slide, giving every cNvPr a fresh id."""
    new = copy.deepcopy(el)
    for cnv in new.iter('{%s}cNvPr' % NS['p']):
        cnv.set('id', str(next_id[0]))
        next_id[0] += 1
    dst_spTree.append(new)
    return new


R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def carry_picture(pic_el, src_slide, dst_slide):
    """A copied <p:pic> still points at the source slide's relationship id: add the image to the destination slide's
    package and point the blip at the new relationship."""
    import io
    for blip in pic_el.iter('{%s}blip' % NS['a']):
        rid = blip.get('{%s}embed' % R_NS)
        if not rid:
            continue
        image_part = src_slide.part.related_part(rid)
        new_part, new_rid = dst_slide.part.get_or_add_image_part(io.BytesIO(image_part.blob))
        blip.set('{%s}embed' % R_NS, new_rid)


def retype(el, font):
    """Point every explicit typeface at the deck's font (Noto Sans CJK SC exists on the render host, not on the author's PC)."""
    for tag in ('latin', 'ea', 'cs'):
        for f in el.iter('{%s}%s' % (NS['a'], tag)):
            f.set('typeface', font)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--deck', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--model', type=int, default=9, help='1-based slide whose title bar style is copied')
    ap.add_argument('--after', type=int, default=9, help='insert the new slides after this 1-based slide')
    ap.add_argument('--font', default=None, help='retype copied text to this typeface, e.g. 微软雅黑')
    ap.add_argument('items', nargs='+', help='<editable.pptx>::<title>')
    ap.add_argument('--any-height', action='store_true', help='mechanism tests only: accept a scene that is not 16:9')
    a = ap.parse_args()

    prs = Presentation(a.deck)
    model = prs.slides[a.model - 1]
    layout = model.slide_layout
    bars = [sh for sh in model.shapes if re.match(r'矩形 [23]$', sh.name)]
    assert len(bars) == 2, [sh.name for sh in model.shapes][:5]
    sldIdLst = prs.slides._sldIdLst
    insert_at = a.after

    for item in a.items:
        src_path, title = item.split('::', 1)
        src = Presentation(src_path)
        assert abs(src.slide_width - SLIDE_W_EMU) < 2000, f'{src_path}: slide width {src.slide_width} EMU is not 960 pt'
        assert a.any_height or abs(src.slide_height - prs.slide_height) < 200000, \
            f'{src_path}: slide height {src.slide_height / 12700:.0f} pt, deck {prs.slide_height / 12700:.0f} pt (scene must be 1536×864)'
        s_src = src.slides[0]

        slide = prs.slides.add_slide(layout)
        # only the title placeholder survives; the deck's content slides carry no body placeholders
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type != 1:   # 1 = TITLE
                ph._element.getparent().remove(ph._element)
        spTree = slide.shapes._spTree
        next_id = [max(int(c.get('id')) for c in spTree.iter('{%s}cNvPr' % NS['p'])) + 1]
        for b in bars:
            copy_shape(b._element, spTree, next_id)
        # the title: same run properties as the model slide
        t_new = slide.shapes.title
        t_model = copy.deepcopy(title_like(model)._element)
        t_new._element.getparent().replace(t_new._element, t_model)
        t_model.find('p:nvSpPr/p:cNvPr', NS).set('id', str(next_id[0])); next_id[0] += 1
        set_title(t_model, title)
        # the figure: every spTree child that is a shape, 1:1
        n = 0
        for el in s_src.shapes._spTree:
            tag = el.tag.split('}')[1]
            if tag in ('nvGrpSpPr', 'grpSpPr'):
                continue
            new = copy_shape(el, spTree, next_id)
            if tag == 'pic':
                carry_picture(new, s_src, slide)
            if a.font:
                retype(new, a.font)
            n += 1
        # move the new slide to its place
        sid = sldIdLst[-1]
        sldIdLst.remove(sid)
        sldIdLst.insert(insert_at, sid)
        insert_at += 1
        print(f'slide {insert_at}: "{title}" ← {Path(src_path).name}: {n} shapes')

    prs.save(a.out)
    print('wrote', a.out, 'slides', len(Presentation(a.out).slides))


if __name__ == '__main__':
    main()
